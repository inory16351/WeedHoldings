import openpyxl
import json
import sys
import os

def parse_excel(xlsx_path, json_path):
    if not os.path.exists(xlsx_path):
        print(f"Error: {xlsx_path} not found.")
        sys.exit(1)
        
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    # 1. 한번에심기_한번에수확 그룹 sheet to build plant ID -> Group ID mapping
    group_map = {}
    if "한번에심기_한번에수확 그룹" in wb.sheetnames:
        sheet = wb["한번에심기_한번에수확 그룹"]
        for r in range(4, sheet.max_row + 1):
            p_id = sheet.cell(row=r, column=3).value
            g_id = sheet.cell(row=r, column=2).value
            if p_id is not None and g_id is not None:
                group_map[int(p_id)] = int(g_id)
                
    # 2. 원료 식물 sheet
    plants = []
    if "원료 식물" in wb.sheetnames:
        sheet = wb["원료 식물"]
        for r in range(4, sheet.max_row + 1):
            p_id = sheet.cell(row=r, column=1).value
            name = sheet.cell(row=r, column=2).value
            if p_id is None or name is None:
                continue
            
            p_id = int(p_id)
            grow_time = float(sheet.cell(row=r, column=3).value or 0.0)
            req_level = int(sheet.cell(row=r, column=4).value or 0)
            req_gold = int(sheet.cell(row=r, column=5).value or 0)
            icon = sheet.cell(row=r, column=6).value or ""
            baby = sheet.cell(row=r, column=7).value or ""
            full_grow = sheet.cell(row=r, column=8).value or ""
            water_time = float(sheet.cell(row=r, column=9).value or 0.0)
            
            group_id = group_map.get(p_id, 0)
            
            plants.append({
                "id": p_id,
                "name": name,
                "growTime": grow_time,
                "reqLevel": req_level,
                "reqGold": req_gold,
                "icon": icon,
                "baby": baby,
                "fullGrow": full_grow,
                "waterTime": water_time,
                "groupId": group_id
            })
            
    # 3. 시설 업그레이드 sheet
    upgrades = []
    if "시설 업그레이드" in wb.sheetnames:
        sheet = wb["시설 업그레이드"]
        for r in range(4, sheet.max_row + 1):
            up_id = sheet.cell(row=r, column=1).value
            level = sheet.cell(row=r, column=2).value
            if up_id is None or level is None:
                continue
                
            up_id = int(up_id)
            level = int(level)
            grow_bonus = float(sheet.cell(row=r, column=3).value or 0.0)
            req_gold = int(sheet.cell(row=r, column=4).value or 0)
            unlock_field = sheet.cell(row=r, column=5).value
            unlock_field = int(unlock_field) if unlock_field is not None else 0
            harvest_group_id = sheet.cell(row=r, column=6).value
            harvest_group_id = int(harvest_group_id) if harvest_group_id is not None else 0
            
            upgrades.append({
                "upgradeID": up_id,
                "level": level,
                "growBonus": grow_bonus,
                "reqGold": req_gold,
                "unlockField": unlock_field,
                "harvestGroupID": harvest_group_id
            })
            
    result = {
        "plants": plants,
        "upgrades": upgrades
    }
    
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"Successfully wrote JSON to {json_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python parse_excel.py <xlsx_path> <output_json_path>")
        sys.exit(1)
    parse_excel(sys.argv[1], sys.argv[2])
