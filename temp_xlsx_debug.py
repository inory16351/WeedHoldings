import openpyxl
import os

xlsx_path = r"c:\Project\슬기로운재배생활\슬기로운재배생활_재배테이블.xlsx"

if not os.path.exists(xlsx_path):
    print("Excel file not found!")
    exit()

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
print("Sheets in workbook:", wb.sheetnames)

for sheetname in wb.sheetnames:
    print(f"\n--- Sheet: {sheetname} ---")
    sheet = wb[sheetname]
    for r in range(1, 15):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 15)]
        # 값들 중 하나라도 있으면 출력
        if any(v is not None for v in row_vals):
            print(f"Row {r:02d}: {row_vals}")

print("\nDump complete.")
