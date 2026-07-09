import openpyxl
import json
import os

xlsx_path = r"c:\Project\슬기로운재배생활\슬기로운재배생활_재배테이블.xlsx"

if not os.path.exists(xlsx_path):
    print("Excel file not found!")
    exit()

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
data = {}

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    rows_data = []
    # Read all rows
    max_row = sheet.max_row
    max_col = sheet.max_column
    for r in range(1, max_row + 1):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        # Filter rows that are completely empty
        if any(v is not None for v in row_vals):
            rows_data.append(row_vals)
    data[sheetname] = rows_data

with open("xlsx_dump_utf8.json", "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("Dumped all sheets to xlsx_dump_utf8.json successfully.")
